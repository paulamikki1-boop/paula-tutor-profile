import re, pathlib
from docx import Document
from docx.shared import Pt, RGBColor, Mm, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image
from ir import all_topics, HOWTO, COVER_SUB, BACK, FOOTER
NAVY, BLUE, MID, ACC, SKY = "0a1628", "1e3a5f", "2a5080", "4a90d9", "e8f0fe"
ORANGE, ORANGE_T, PURPLE = "fdf1e3", "b0651a", "6b4c9a"
rgb = lambda h: RGBColor.from_string(h.upper())
def shade(el, fill):
    pr = el.get_or_add_tcPr() if el.tag.endswith("}tc") else el.get_or_add_pPr()
    s = OxmlElement("w:shd"); s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), fill); pr.append(s)
def cell_borders(cell, color=None, left=None, sz=4):
    tcPr = cell._tc.get_or_add_tcPr(); b = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{side}")
        if side == "left" and left: e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "24"); e.set(qn("w:color"), left)
        elif color: e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz)); e.set(qn("w:color"), color)
        else: e.set(qn("w:val"), "nil")
        b.append(e)
    tcPr.append(b)
def table_borders(t, color="dfe6ef"):
    tblPr = t._tbl.tblPr; b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{side}")
        if side in ("insideH", "bottom"): e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), color)
        else: e.set(qn("w:val"), "nil")
        b.append(e)
    tblPr.append(b)
def repeat_header(row):
    trPr = row._tr.get_or_add_trPr(); e = OxmlElement("w:tblHeader"); e.set(qn("w:val"), "true"); trPr.append(e)


ORDER = {
 "pPr": "pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd tabs suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi adjustRightInd snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc textDirection textAlignment textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange",
 "rPr": "rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss imprint noProof snapToGrid vanish webHidden color spacing w kern position sz szCs highlight u effect bdr shd fitText vertAlign rtl cs em lang eastAsianLayout specVanish oMath",
 "tcPr": "cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar textDirection tcFitText vAlign hideMark",
 "tblPr": "tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize tblW jc tblCellSpacing tblInd tblBorders shd tblLayout tblCellMar tblLook",
}
ORDER = {k: {n: i for i, n in enumerate(v.split())} for k, v in ORDER.items()}
def reorder(root):
    for tag, rank in ORDER.items():
        for el in root.iter(qn("w:" + tag)):
            kids = list(el)
            for k in kids:
                if k.tag.split("}")[1] not in rank: raise SystemExit(f"unknown {k.tag} in {tag}")
            kids.sort(key=lambda k: rank[k.tag.split("}")[1]])
            for k in kids: el.remove(k)
            for k in kids: el.append(k)
def save(d, path):
    reorder(d.element); reorder(d.styles.element)
    for sec in d.sections: reorder(sec.footer._element)
    d.save(path)
BM = [0]
def bookmark(p, name):
    BM[0] += 1
    a = OxmlElement("w:bookmarkStart"); a.set(qn("w:id"), str(BM[0])); a.set(qn("w:name"), name)
    b = OxmlElement("w:bookmarkEnd"); b.set(qn("w:id"), str(BM[0]))
    p._p.insert(1 if p._p.pPr is not None else 0, a); p._p.append(b)
def link(p, text, anchor, size=11):
    h = OxmlElement("w:hyperlink"); h.set(qn("w:anchor"), anchor); h.set(qn("w:history"), "1")
    r = OxmlElement("w:r"); rp = OxmlElement("w:rPr")
    c = OxmlElement("w:color"); c.set(qn("w:val"), "1F2A37"); rp.append(c)
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(size*2)); rp.append(sz)
    r.append(rp); t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve"); r.append(t); h.append(r); p._p.append(h)

class W:
    def __init__(self, scale):
        self.s = scale; self.d = Document(); d = self.d
        sec = d.sections[0]; sec.page_width, sec.page_height = Mm(210), Mm(297)
        m = Mm(16 if scale > 1 else 14); sec.left_margin = sec.right_margin = m; sec.top_margin = Mm(14); sec.bottom_margin = Mm(16)
        st = d.styles
        n = st["Normal"]; n.font.name = "DM Sans"; n.font.size = Pt(10 * scale); n.font.color.rgb = rgb("1f2a37")
        n.element.rPr.rFonts.set(qn("w:eastAsia"), "DM Sans")
        n.paragraph_format.space_after = Pt(4 * scale); n.paragraph_format.line_spacing = 1.15 if scale == 1 else 1.3
        for name, size, col, before in (("Title", 34, NAVY, 0), ("Heading 1", 20, NAVY, 0), ("Heading 2", 13, BLUE, 12), ("Heading 3", 10, MID, 8)):
            h = st[name]; h.font.name = "DM Sans"; h.font.size = Pt(size * scale if name != "Title" else size); h.font.bold = True
            h.font.color.rgb = rgb(col); h.font.italic = False
            rpr = h.element.get_or_add_rPr(); rf = rpr.find(qn("w:rFonts"))
            if rf is None: rf = OxmlElement("w:rFonts"); rpr.append(rf)
            for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"): rf.set(qn(a), "DM Sans")
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                if rf.get(qn(a)) is not None: del rf.attrib[qn(a)]
            h.paragraph_format.space_before = Pt(before * scale); h.paragraph_format.space_after = Pt(4 * scale)
            h.paragraph_format.keep_with_next = True
        st["Heading 3"].font.all_caps = True
        for name in ("List Bullet", "List Number"): st[name].font.name = "DM Sans"
        self.footer()
    def footer(self, sec=None):
        sec = sec or self.d.sections[0]
        p = sec.footer.paragraphs[0]; p.text = ""; r = p.add_run(FOOTER); r.font.size = Pt(7.5); r.font.color.rgb = rgb("8a96a8")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    def runs(self, p, runs, size=None, color=None, bold=None):
        for t, f in runs:
            if t == "\n": p.add_run().add_break(); continue
            if "ext" in f:
                r = p.add_run(" SUPPLEMENT "); r.bold = True; r.font.size = Pt(7 * self.s); r.font.color.rgb = rgb(PURPLE)
                rp = r._r.get_or_add_rPr(); s = OxmlElement("w:shd"); s.set(qn("w:val"), "clear"); s.set(qn("w:fill"), "f5f0fa"); rp.append(s)
                p.add_run(" "); continue
            r = p.add_run(t)
            if "b" in f or bold: r.bold = True
            if "b" in f and not color: r.font.color.rgb = rgb(NAVY)
            if "i" in f: r.italic = True
            if "sub" in f: r.font.subscript = True
            if "sup" in f: r.font.superscript = True
            if size: r.font.size = Pt(size)
            if color: r.font.color.rgb = rgb(color)
        return p
    def para(self, runs, style=None, **kw):
        p = self.d.add_paragraph(style=style); return self.runs(p, runs, **kw)
    def box(self, fill, left, label, labcol, body_fn):
        t = self.d.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = t.cell(0, 0); shade(c._tc, fill); cell_borders(c, color="dfe6ef" if not left else None, left=left)
        p = c.paragraphs[0]; r = p.add_run(label.upper()); r.bold = True; r.font.size = Pt(7.5 * self.s); r.font.color.rgb = rgb(labcol)
        body_fn(c); self.d.add_paragraph().paragraph_format.space_after = Pt(0)
    def blocks(self, blocks):
        d = self.d; s = self.s
        for b in blocks:
            k = b[0]
            if k == "h":
                p = d.add_paragraph(style="Heading 2" if b[1] == 3 else "Heading 3")
                if b[3]: self.runs(p, [("Supplement", frozenset({"ext"}))])
                self.runs(p, b[2])
            elif k == "p": self.para(b[1])
            elif k in ("ul", "ol"):
                for r in b[1]: self.para(r, style="List Bullet" if k == "ul" else "List Number")
            elif k == "table":
                t = d.add_table(rows=1 + len(b[2]), cols=len(b[1])); table_borders(t)
                for j, h in enumerate(b[1]):
                    c = t.cell(0, j); shade(c._tc, SKY); p = c.paragraphs[0]
                    self.runs(p, [(x.upper() if x != "\n" else x, f) for x, f in h], size=7.5 * s, color=MID, bold=True)
                repeat_header(t.rows[0])
                for i, row in enumerate(b[2]):
                    for j, cl in enumerate(row[:len(b[1])]): self.runs(t.cell(i + 1, j).paragraphs[0], cl, size=9 * s)
                d.add_paragraph().paragraph_format.space_after = Pt(0)
            elif k == "conv":
                t = d.add_table(rows=len(b[1]), cols=2); table_borders(t)
                for i, (lab, v) in enumerate(b[1]):
                    self.runs(t.cell(i, 0).paragraphs[0], [(lab.upper(), frozenset())], size=7.5 * s, color=ACC, bold=True)
                    self.runs(t.cell(i, 1).paragraphs[0], v, size=9.3 * s)
                    t.cell(i, 0).width = Mm(45); t.cell(i, 1).width = Mm(135)
                d.add_paragraph().paragraph_format.space_after = Pt(0)
            elif k == "steps":
                for num, title, body in b[1]:
                    p = d.add_paragraph(); p.paragraph_format.left_indent = Mm(7); p.paragraph_format.first_line_indent = Mm(-7)
                    r = p.add_run(f"{num}\t"); r.bold = True; r.font.color.rgb = rgb(ACC)
                    self.runs(p, title, bold=True, color=NAVY); p.add_run().add_break()
                    self.runs(p, body, color="5b6778")
            elif k == "mnemonic":
                p = self.para(b[1]); shade(p._p, SKY)
            elif k == "box":
                kind = b[1]
                fill, left, labcol = {"tip": (SKY, ACC, MID), "mistake": (ORANGE, "d9822b", ORANGE_T), "example": ("fbfcfe", None, MID)}[kind]
                self.box(fill, left, b[2], labcol, lambda c, r=b[3]: self.runs(c.add_paragraph(), r, size=9.4 * s))
            elif k == "checklist":
                def body(c, items=b[2]):
                    for it in items: p = c.add_paragraph(); r = p.add_run("☐  "); r.font.color.rgb = rgb(ACC); self.runs(p, it, size=9.3 * s)
                self.box("ffffff", None, b[1], ACC, body)
            elif k == "figure":
                img = pathlib.Path("assets-chem")/f"{b[1]}.png"; w, h = Image.open(img).size
                width = Mm(130 if s == 1 else 150); p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(img), width=width); p.paragraph_format.keep_with_next = True
                c = self.para([(b[2], frozenset({"i"}))], color="5b6778", size=9 * s); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    def topic_head(self, meta, title):
        d = self.d
        p = d.add_paragraph(); r = p.add_run(meta["badge"].upper()); r.bold = True; r.font.size = Pt(8.5 * self.s); r.font.color.rgb = rgb(ACC)
        p.paragraph_format.space_after = Pt(0); p.paragraph_format.keep_with_next = True
        h1 = d.add_paragraph(title, style="Heading 1"); bookmark(h1, "topic%02d" % meta["n"])
        self.para(meta["purpose"], color="5b6778")
    def cover(self, top, kicker, title, sub, extra=None):
        d = self.d
        t = d.add_table(rows=1, cols=1); c = t.cell(0, 0); shade(c._tc, NAVY); cell_borders(c)
        trPr = t.rows[0]._tr.get_or_add_trPr(); hh = OxmlElement("w:trHeight"); hh.set(qn("w:val"), str(int(Mm(250) / 635))); hh.set(qn("w:hRule"), "exact"); trPr.append(hh)
        def line(text, size, col, bold=False, italic=False, before=0, font=None):
            p = c.add_paragraph(); r = p.add_run(text); r.font.size = Pt(size); r.font.color.rgb = rgb(col); r.bold = bold; r.italic = italic
            if font: r.font.name = font
            p.paragraph_format.space_before = Pt(before); p.paragraph_format.left_indent = Mm(10); p.paragraph_format.right_indent = Mm(10); return p
        c.paragraphs[0].paragraph_format.space_before = Pt(20)
        line(top.upper(), 10, "9fc3ee", bold=True, before=10)
        line(kicker, 11, "9fc3ee", before=170)
        line(title[0], 48, "ffffff", bold=True, font="Cormorant Garamond", before=6)
        if title[1]: line(title[1], 28, ACC, bold=True, italic=True, font="Cormorant Garamond")
        line(sub, 12, "dbe7f7", before=10)
        if extra: line(extra, 10, "dbe7f7", before=14)
        line("Paula.", 24, "ffffff", bold=True, font="Cormorant Garamond", before=60)
        line("IGCSE, GCSE & A Level Tutor · Kuala Lumpur & Online · paulamikki1-boop.github.io/paula-tutor-profile · WhatsApp +60 10-368 6902", 9, "9fc3ee")
        d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    def toc_field(self):
        p = self.d.add_paragraph(); r = p.add_run()
        for tag, txt in (("begin", None), (None, 'TOC \\o "1-1" \\h \\z \\u'), ("separate", None)):
            if tag: e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), tag); r._r.append(e)
            else: e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt; r._r.append(e)
        return p, r
    def pagebreak(self): self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

def endfield(r):
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end"); r._r.append(e)
def settings_update(d):
    e = OxmlElement("w:updateFields"); e.set(qn("w:val"), "true"); d.settings.element.append(e)

T = all_topics()
for n, *_rest, meta, blocks in T: meta['n'] = n
out = pathlib.Path("out/Word"); (out/"Detailed Notes (by topic)").mkdir(parents=True, exist_ok=True)
# Compact
w = W(1); d = w.d
d.core_properties.title = "IGCSE Chemistry Revision Notes (Cambridge 0620, 2026-2028)"; d.core_properties.author = "Paula"
w.cover("IGCSE Revision Notes", "Cambridge IGCSE 0620 · 2026–2028 syllabus", ("Chemistry", "Revision Notes"), COVER_SUB,
        "SUPPLEMENT = Extended tier only (Papers 2 & 4) · blue boxes = examiner tips · orange boxes = common mistakes · ☐ = tick-off checklist")
p = d.add_paragraph(); r = p.add_run("Contents"); r.bold = True; r.font.size = Pt(22); r.font.color.rgb = rgb(NAVY)
grp = None
for n, s_, t, g, meta, blocks in T:
    if g != grp:
        gp = d.add_paragraph(); gr = gp.add_run(g.upper()); gr.bold = True; gr.font.size = Pt(8); gr.font.color.rgb = rgb(ACC)
        gp.paragraph_format.space_before = Pt(8); gp.paragraph_format.space_after = Pt(1); grp = g
    lp = d.add_paragraph(); lp.paragraph_format.space_after = Pt(2)
    nr = lp.add_run(f"{n}   "); nr.bold = True; nr.font.color.rgb = rgb(MID)
    link(lp, t, "topic%02d" % n)
d.add_paragraph("How to use these notes", style="Heading 2")
for h in HOWTO: d.add_paragraph(h, style="List Number")
for n, s_, t, g, meta, blocks in T:
    w.pagebreak(); w.topic_head(meta, t); w.blocks(blocks)
w.pagebreak()
for i, line in enumerate(BACK):
    pp = d.add_paragraph(); rr = pp.add_run(line); pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr.font.size = Pt(22 if i == 0 else 12); rr.bold = i in (0, 2); rr.font.color.rgb = rgb(NAVY if i == 0 else MID if i == 2 else "5b6778")
save(d, out/"IGCSE-Chemistry-Compact-Revision-Notes.docx")
# Detailed
for n, s, t, g, meta, blocks in T:
    w = W(1.15); d = w.d
    d.core_properties.title = f"IGCSE Chemistry Detailed Notes: Topic {n} {t}"; d.core_properties.author = "Paula"
    w.cover("IGCSE Chemistry · Detailed Notes", "Cambridge IGCSE 0620 · 2026–2028 syllabus", (f"{n:02d}", t), "".join(x for x, f in meta["purpose"] if x != "\n"))
    w.topic_head(meta, t); w.blocks(blocks)
    pp = w.para([("Want your answers checked? ", frozenset({"b"})), ("Send me a past paper and I’ll mark it with examiner-style feedback within 48 hours: paulamikki1-boop.github.io/paula-tutor-profile/marking · WhatsApp +60 10-368 6902", frozenset())])
    shade(pp._p, SKY)
    name = re.sub(r"[^A-Za-z0-9]+", "-", t).strip("-")
    save(d, out/"Detailed Notes (by topic)"/f"{n:02d}-{name}.docx")
print("ok")
